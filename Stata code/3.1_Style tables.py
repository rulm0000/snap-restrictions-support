"""Apply Allcott/booktabs-style horizontal rules and merges to SNAP Excel tables."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

TABLES = Path(sys.argv[1]) / "Tables"

THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")
DOUBLE = Side(style="double", color="000000")
NO_SIDE = Side(style=None)

SECTION_PREFIXES = (
    "Psychological factors",
    "Psychological predictors",  # legacy spelling, pre-8/17 xlsx
    "Covariates",
    "Total sample",
)


def _clear_borders(cell) -> None:
    cell.border = Border(left=NO_SIDE, right=NO_SIDE, top=NO_SIDE, bottom=NO_SIDE)


def _set_bottom(cell, side: Side) -> None:
    b = cell.border
    cell.border = Border(left=b.left, right=b.right, top=b.top, bottom=side)


def _set_top(cell, side: Side) -> None:
    b = cell.border
    cell.border = Border(left=b.left, right=b.right, top=side, bottom=b.bottom)


def _set_left(cell, side: Side) -> None:
    b = cell.border
    cell.border = Border(left=side, right=b.right, top=b.top, bottom=b.bottom)


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_empty(value) -> bool:
    return _cell_text(value) == ""


def _last_used_row(ws) -> int:
    last = 1
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        if any(not _is_empty(c.value) for c in row):
            last = row[0].row
    return last


def _last_used_col(ws, max_row: int) -> int:
    last = 1
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=ws.max_column):
        for c in row:
            if not _is_empty(c.value):
                last = max(last, c.column)
    return last


def _is_section_header(value) -> bool:
    text = _cell_text(value)
    if not text:
        return False
    return any(text.startswith(p) for p in SECTION_PREFIXES)


def _unmerge_all(ws) -> None:
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))


def _safe_merge(ws, start_row: int, start_col: int, end_row: int, end_col: int) -> None:
    if end_row < start_row or end_col < start_col:
        return
    if start_row == end_row and start_col == end_col:
        return
    try:
        ws.merge_cells(
            start_row=start_row,
            start_column=start_col,
            end_row=end_row,
            end_column=end_col,
        )
    except ValueError:
        pass


def _row_is_label_only(ws, row: int, last_col: int) -> bool:
    """True when column A has a label and all other columns are empty (domain/section)."""
    if _is_empty(ws.cell(row, 1).value):
        return False
    return all(_is_empty(ws.cell(row, c).value) for c in range(2, last_col + 1))


def _merge_spans(ws, last_row: int, last_col: int) -> None:
    # Title across full width
    _safe_merge(ws, 1, 1, 1, last_col)
    title = ws.cell(1, 1)
    title.font = Font(name=title.font.name or "Calibri", size=11, bold=True)
    title.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Section / domain headers: label in A, empty elsewhere
    for r in range(2, last_row + 1):
        if _row_is_label_only(ws, r, last_col):
            _safe_merge(ws, r, 1, r, last_col)
            cell = ws.cell(r, 1)
            # Keep bold for section banners; domain names stay regular unless already bold
            if _is_section_header(cell.value):
                cell.font = Font(
                    name=cell.font.name or "Calibri",
                    size=cell.font.size or 11,
                    bold=True,
                    italic=bool(cell.font.italic),
                )
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_sheet(
    ws,
    *,
    header_rows: int = 1,
    colnum_row: int | None = None,
    table3: bool = False,
) -> None:
    last_row = _last_used_row(ws)
    last_col = _last_used_col(ws, last_row)
    if last_row < 2:
        return

    _unmerge_all(ws)

    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=last_col):
        for cell in row:
            _clear_borders(cell)

    top_rule_row = 2
    if colnum_row is not None:
        header_bottom = colnum_row
    else:
        header_bottom = 1 + header_rows

    for col in range(1, last_col + 1):
        _set_top(ws.cell(top_rule_row, col), DOUBLE)
        _set_bottom(ws.cell(header_bottom, col), THIN)

    for r in range(header_bottom + 1, last_row + 1):
        if _is_section_header(ws.cell(r, 1).value):
            for col in range(1, last_col + 1):
                _set_top(ws.cell(r, col), THIN)

    for col in range(1, last_col + 1):
        _set_bottom(ws.cell(last_row, col), DOUBLE)

    # Table 3: group headers + vertical separators between SNAP blocks
    if table3 and last_col >= 7:
        # Merge Non-recipients (B–D) and SNAP recipients (E–G)
        _safe_merge(ws, 2, 2, 2, 4)
        _safe_merge(ws, 2, 5, 2, 7)
        # Span p-for-interaction over the two header rows
        _safe_merge(ws, 2, 8, 3, 8)

        for col in range(2, 5):
            _set_bottom(ws.cell(2, col), THIN)
        for col in range(5, 8):
            _set_bottom(ws.cell(2, col), THIN)

        # Vertical rules separating Non-recipients | SNAP recipients | interaction p
        for r in range(2, last_row + 1):
            _set_left(ws.cell(r, 5), MEDIUM)  # before SNAP recipients
            _set_left(ws.cell(r, 8), MEDIUM)  # before p for interaction

        for col in (2, 5, 8):
            ws.cell(2, col).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        ws.cell(2, 8).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    # Alignment
    for r in range(1, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            if c == 1 and r > 1:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif r > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    _merge_spans(ws, last_row, last_col)

    # Re-apply centered alignment on merged Table 3 group headers
    if table3 and last_col >= 7:
        for col in (2, 5, 8):
            ws.cell(2, col).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

    ws.column_dimensions["A"].width = 42
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16


def style_table1(path: Path) -> None:
    wb = load_workbook(path)
    style_sheet(wb.active, header_rows=1, colnum_row=None)
    wb.save(path)
    print(f"styled {path.name}")


def style_table1a(path: Path) -> None:
    wb = load_workbook(path)
    style_sheet(wb.active, header_rows=1, colnum_row=None)
    wb.save(path)
    print(f"styled {path.name}")


def style_table2(path: Path) -> None:
    wb = load_workbook(path)
    style_sheet(wb.active, header_rows=1, colnum_row=None)
    wb.save(path)
    print(f"styled {path.name}")


def style_table3(path: Path) -> None:
    wb = load_workbook(path)
    # Two header rows (group + β/d/p)
    style_sheet(wb.active, header_rows=2, colnum_row=None, table3=True)
    wb.save(path)
    print(f"styled {path.name}")


def main() -> None:
    jobs = [
        (TABLES / "Table1_Descriptives.xlsx", style_table1),
        (TABLES / "Table1a_Percent_Support.xlsx", style_table1a),
        (TABLES / "Table2_Regression.xlsx", style_table2),
        (TABLES / "Table3_Predictors_by_SNAP.xlsx", style_table3),
    ]
    for path, fn in jobs:
        if not path.exists():
            print(f"missing {path.name}; skip")
            continue
        fn(path)


if __name__ == "__main__":
    main()
